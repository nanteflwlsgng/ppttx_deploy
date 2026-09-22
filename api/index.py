from http.server import BaseHTTPRequestHandler, HTTPServer
import json
import os
import io
from pptx import Presentation
from pptx.util import Pt

class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        # 1. Lire la longueur du contenu
        content_length = int(self.headers.get('Content-Length', 0))
        post_data = self.rfile.read(content_length)
        
        try:
            data = json.loads(post_data.decode('utf-8'))
        except Exception as e:
            self.send_response(400)
            self.end_headers()
            self.wfile.write(str(e).encode())
            return

        # 2. Charger le template
        template_path = os.path.join(os.path.dirname(__file__), 'template.pptx')

        try:
            if os.path.exists(template_path):
                prs = Presentation(template_path)
            else:
                prs = Presentation()
                prs.slide_width = 12192000
                prs.slide_height = 6858000
        except Exception:
            prs = Presentation()
            prs.slide_width = 12192000
            prs.slide_height = 6858000

        # 3. Remplir les slides
        slides_content = data.get('slides', [])
        for slide_data in slides_content:
            try:
                layout = prs.slide_layouts[3] # Deux contenus
            except:
                layout = prs.slide_layouts[1]

            slide = prs.slides.add_slide(layout)
            
            # Titre
            if slide.shapes.title:
                slide.shapes.title.text = slide_data.get('titre', 'Sans titre')

            # Contenu Gauche
            points = slide_data.get('points', [])
            if len(slide.placeholders) > 1:
                left_shape = slide.placeholders[1]
                tf = left_shape.text_frame
                tf.text = ""
                tf.word_wrap = True
                for point in points:
                    p = tf.add_paragraph()
                    p.text = point
                    p.level = 0
                    p.font.size = Pt(16)
                    p.space_after = Pt(10)

        # 4. Sauvegarder en mémoire (RAM)
        ppt_stream = io.BytesIO()
        prs.save(ppt_stream)
        ppt_stream.seek(0)

        # 5. Envoyer la réponse HTTP
        self.send_response(200)
        self.send_header('Content-type', 'application/vnd.openxmlformats-officedocument.presentationml.presentation')
        self.send_header('Content-Disposition', 'attachment; filename="soutenance.pptx"')
        self.end_headers()
        self.wfile.write(ppt_stream.read())
        return

#local python
if __name__ == '__main__':
    server_address = ('127.0.0.1', 5328)
    httpd = HTTPServer(server_address, handler)
    print("🐍 Serveur Python GhostPPTX actif sur http://127.0.0.1:5328", flush=True)
    httpd.serve_forever()