import sys
import json
import os
import io
from pptx import Presentation
from pptx.util import Pt # Nécessaire pour contrôler la taille du texte

# 1. FORCER L'ENCODAGE UTF-8
sys.stdin = io.TextIOWrapper(sys.stdin.buffer, encoding='utf-8')

# 2. LIRE LES DONNÉES
try:
    input_data = sys.stdin.read()
    data = json.loads(input_data)
except Exception as e:
    sys.stderr.write(f"Erreur lecture JSON: {str(e)}")
    sys.exit(1)

# 3. CHARGER LE TEMPLATE
template_path = os.path.join(os.getcwd(), 'template.pptx')

try:
    if os.path.exists(template_path):
        prs = Presentation(template_path)
    else:
        sys.stderr.write("ATTENTION: Template introuvable. Création d'un vide.\n")
        prs = Presentation()
        prs.slide_width = 12192000 
        prs.slide_height = 6858000
except Exception as e:
    sys.stderr.write(f"Erreur chargement PPTX: {str(e)}")
    sys.exit(1)

# 4. REMPLIR LES SLIDES
slides_content = data.get('slides', [])

for slide_data in slides_content:
    # --- CHOIX DU LAYOUT (Index 3 = Deux Contenus) ---
    try:
        # On utilise le layout qui a 2 colonnes
        layout = prs.slide_layouts[3] 
    except:
        layout = prs.slide_layouts[1] # Fallback

    slide = prs.slides.add_slide(layout)

    # --- TITRE ---
    if slide.shapes.title:
        slide.shapes.title.text = slide_data.get('titre', 'Sans titre')

    # --- REMPLISSAGE COLONNE GAUCHE UNIQUEMENT ---
    points = slide_data.get('points', [])
    
    # Vérification : Est-ce qu'on a bien la zone de gauche ?
    # Généralement : placeholders[0]=Titre, placeholders[1]=Gauche, placeholders[2]=Droite
    if len(slide.placeholders) > 1:
        left_shape = slide.placeholders[1]
        
        tf = left_shape.text_frame
        tf.text = "" # On vide le texte par défaut ("Cliquez pour...")
        tf.word_wrap = True # Active le retour à la ligne automatique

        for point in points:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0 # Niveau de puce principal
            
            # --- ASTUCE ANTI-DÉBORDEMENT ---
            # On fixe la police à 18 points (Standard pro). 
            # Si le texte est très long, tu peux descendre à 16 ou 14.
            p.font.size = Pt(18) 
            
            # Un peu d'espace entre les points pour la lisibilité
            p.space_after = Pt(10)

    # --- COLONNE DROITE ---
    # ON NE FAIT RIEN ICI ! 
    # En ne touchant pas à slide.placeholders[2], PowerPoint garde
    # la zone pointillée "Cliquez pour ajouter du texte / Image / Graphique".
    # C'est parfait pour l'étudiant.

# 5. SAUVEGARDER
output_filename = "temp_output.pptx"
prs.save(output_filename)
print(output_filename)