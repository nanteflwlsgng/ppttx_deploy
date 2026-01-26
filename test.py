import sys
import json
import os
import io
import math
from pptx import Presentation

# 1. FORCER L'ENCODAGE UTF-8 (Vital pour les accents)
sys.stdin = io.TextIOWrapper(sys.stdin.buffer, encoding='utf-8')

# 2. LIRE LES DONNÉES
try:
    input_data = sys.stdin.read()
    data = json.loads(input_data)
except Exception as e:
    sys.stderr.write(f"Erreur lecture JSON: {str(e)}")
    sys.exit(1)

# 3. CHARGER LE TEMPLATE "INTEGRAL"
template_path = os.path.join(os.getcwd(), 'template.pptx')

try:
    if os.path.exists(template_path):
        prs = Presentation(template_path)
    else:
        # Fallback si le template n'est pas là
        sys.stderr.write("ATTENTION: Template introuvable. Création d'un vide.\n")
        prs = Presentation()
        prs.slide_width = 12192000  # Force 16:9 en EMUs
        prs.slide_height = 6858000
except Exception as e:
    sys.stderr.write(f"Erreur chargement PPTX: {str(e)}")
    sys.exit(1)

# 4. REMPLIR LES SLIDES
slides_content = data.get('slides', [])

for slide_data in slides_content:
    # --- CHOIX DU LAYOUT ---
    # Index 3 est SOUVENT "Deux Contenus" dans les thèmes standards.
    # Si ça ne marche pas, essaie 4.
    # layout[0] = Titre, layout[1] = Contenu simple, layout[3] = Deux Contenus
    try:
        layout = prs.slide_layouts[3] 
    except:
        layout = prs.slide_layouts[1] # Fallback sur simple contenu

    slide = prs.slides.add_slide(layout)

    # --- TITRE ---
    if slide.shapes.title:
        slide.shapes.title.text = slide_data.get('titre', 'Sans titre')

    # --- GESTION DES DEUX COLONNES ---
    points = slide_data.get('points', [])
    
    # On vérifie si on a bien accès aux placeholders gauche et droite
    # Dans "Deux Contenus", placeholder[0]=Titre, placeholder[1]=Gauche, placeholder[2]=Droite
    if len(slide.placeholders) > 2:
        left_shape = slide.placeholders[1]
        right_shape = slide.placeholders[2]
        
        # On divise les points en deux listes
        half = math.ceil(len(points) / 2)
        left_points = points[:half]
        right_points = points[half:]

        # Remplir Gauche
        tf_left = left_shape.text_frame
        tf_left.text = "" # Vider le texte par défaut du template
        for p_text in left_points:
            p = tf_left.add_paragraph()
            p.text = p_text
            p.level = 0
            # Petite astuce : espacement pour aérer
            p.space_after = 100000 

        # Remplir Droite
        tf_right = right_shape.text_frame
        tf_right.text = ""
        for p_text in right_points:
            p = tf_right.add_paragraph()
            p.text = p_text
            p.level = 0
            p.space_after = 100000

    # --- FALLBACK (Si le layout n'a qu'une seule zone de texte) ---
    elif len(slide.placeholders) > 1:
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.text = ""
        for p_text in points:
            p = tf.add_paragraph()
            p.text = p_text
            p.level = 0

# 5. SAUVEGARDER ET ENVOYER LE NOM DU FICHIER
output_filename = "temp_output.pptx"
prs.save(output_filename)
print(output_filename)